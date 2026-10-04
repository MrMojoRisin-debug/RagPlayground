"""Turn corpus files into text segments.

Each loader returns a list of (text, extra_metadata) segments. Most formats are a
single segment; PDFs yield one segment per page so citations can show page numbers.
"""

from pathlib import Path

Segment = tuple[str, dict]


def _load_text(path: Path) -> list[Segment]:
    return [(path.read_text(encoding="utf-8", errors="replace"), {})]


def _load_pdf(path: Path) -> list[Segment]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [
        (page.extract_text() or "", {"page": i + 1})
        for i, page in enumerate(reader.pages)
    ]


def _load_docx(path: Path) -> list[Segment]:
    import docx

    doc = docx.Document(str(path))
    lines = []
    for para in doc.paragraphs:
        style = (para.style.name or "").lower() if para.style else ""
        text = para.text.strip()
        if not text:
            continue
        # Keep headings recognisable so the chunker can split on them.
        if style.startswith("heading"):
            level = "".join(c for c in style if c.isdigit()) or "1"
            text = "#" * int(level) + " " + text
        lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [("\n\n".join(lines), {})]


def _load_html(path: Path) -> list[Segment]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header"]):
        tag.decompose()
    return [(soup.get_text("\n"), {})]


LOADERS = {
    ".md": _load_text,
    ".markdown": _load_text,
    ".txt": _load_text,
    ".rst": _load_text,
    ".csv": _load_text,
    ".json": _load_text,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".html": _load_html,
    ".htm": _load_html,
}

SUPPORTED_EXTENSIONS = set(LOADERS)


def load_file(path: Path) -> list[Segment]:
    return LOADERS[path.suffix.lower()](path)
